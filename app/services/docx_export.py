import json
import os, re, secrets
from docx import Document as DocxDocument
from docx.shared import Inches

def _sanitize_filename(name: str) -> str:
    name = re.sub(r"[^a-zA-Z0-9._-]+", "_", name.strip())[:180]
    return name or "artifact"

def _load_brand_header() -> dict:
    path = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'data', 'brand_header.json'))
    if not os.path.exists(path):
        return {"brand": "", "email": "", "phone": ""}
    try:
        with open(path, "r", encoding="utf-8") as f:
            data = json.load(f)
        return {
            "brand": (data.get("brand") or "").strip(),
            "email": (data.get("email") or "").strip(),
            "phone": (data.get("phone") or "").strip(),
        }
    except Exception:
        return {"brand": "", "email": "", "phone": ""}

def export_docx(storage_root: str, org_id: int, client_id: int, title: str, body: str) -> str:
    # Save docx under org/client exports
    folder = os.path.join(storage_root, f"org_{org_id}", f"client_{client_id}", "_exports")
    os.makedirs(folder, exist_ok=True)
    token = secrets.token_hex(6)
    fname = f"{token}__{_sanitize_filename(title)}.docx"
    path = os.path.join(folder, fname)

    doc = DocxDocument()

    brand = _load_brand_header()
    # Header with logo + brand (optional)
    try:
        logo = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'static', 'brand', 'logo.png'))
        if os.path.exists(logo):
            header = doc.sections[0].header
            header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            run = header_para.add_run()
            run.add_picture(logo, width=Inches(1.2))
            if brand.get("brand"):
                header_para.add_run("  " + brand.get("brand"))
        elif brand.get("brand"):
            header = doc.sections[0].header
            header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            header_para.add_run(brand.get("brand"))
        if brand.get("email") or brand.get("phone"):
            header_info = doc.sections[0].header.add_paragraph()
            header_info.add_run(" | ".join([v for v in [brand.get("email"), brand.get("phone")] if v]))
    except Exception:
        pass

    doc.add_heading(title or "Document", level=1)
    for line in (body or "").splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    doc.save(path)
    return path


def export_docx_named(storage_root: str, org_id: int, client_id: int, filename: str, title: str, body: str) -> str:
    folder = os.path.join(storage_root, f"org_{org_id}", f"client_{client_id}", "_exports")
    os.makedirs(folder, exist_ok=True)
    filename = _sanitize_filename(filename)
    if not filename.lower().endswith(".docx"):
        filename += ".docx"
    path = os.path.join(folder, filename)

    doc = DocxDocument()

    brand = _load_brand_header()
    # Header with logo + brand (optional)
    try:
        logo = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'static', 'brand', 'logo.png'))
        if os.path.exists(logo):
            header = doc.sections[0].header
            header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            run = header_para.add_run()
            run.add_picture(logo, width=Inches(1.2))
            if brand.get("brand"):
                header_para.add_run("  " + brand.get("brand"))
        elif brand.get("brand"):
            header = doc.sections[0].header
            header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            header_para.add_run(brand.get("brand"))
        if brand.get("email") or brand.get("phone"):
            header_info = doc.sections[0].header.add_paragraph()
            header_info.add_run(" | ".join([v for v in [brand.get("email"), brand.get("phone")] if v]))
    except Exception:
        pass

    doc.add_heading(title or "Document", level=1)
    for line in (body or "").splitlines():
        if line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    doc.save(path)
    return path
